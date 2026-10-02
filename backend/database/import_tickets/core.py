"""Pure mapping and workbook reader. No database connection or side effects."""
from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path


def norm(value):
    text = unicodedata.normalize("NFKD", "" if value is None else str(value)).casefold()
    text = "".join(c for c in text if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def meaningful(value):
    return value is not None and str(value).strip() not in ("", "-", "—")


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False,
                                     default=str).encode("utf-8")).hexdigest()


def load_mapping(path):
    path = Path(path)
    config = json.loads(path.read_text(encoding="utf-8-sig"))
    if config.get("version") != 1:
        raise ValueError("Versão de mapeamento não suportada")
    catalog = json.loads((path.parent / config["catalog_file"]).read_text(encoding="utf-8-sig"))
    if config.get("field_catalog_file"):
        from .field_catalog import read_catalog
        config["field_catalog"] = read_catalog(path.parent / config["field_catalog_file"], config)
    for file_key, value_key in (("resolutions_file", "location_resolutions"),
                                ("summaries_file", "description_summaries"),
                                ("numbering_file", "request_numbers")):
        if config.get(file_key):
            config[value_key] = json.loads((path.parent / config[file_key]).read_text(encoding="utf-8-sig"))
    return config, catalog


def location_signature(cells):
    return digest(sorted((norm(c["header"]), norm(c["value"]), c["role"]) for c in cells
                         if c["role"] in {"business", "region", "location"} and meaningful(c["value"])))


def read_workbooks(paths):
    # Preserve physical columns: repeated/normalized headers must not overwrite values.
    import openpyxl
    from openpyxl.utils import get_column_letter

    for path in paths:
        path = Path(path)
        book = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            for sheet in book:
                rows = sheet.iter_rows()
                first = next(rows, ())
                headers = [str(c.value or "").strip() for c in first]
                if not any(norm(h) == "numero" for h in headers):
                    raise ValueError(f"Cabeçalho Número ausente: {path.name}/{sheet.title}")
                for number, cells in enumerate(rows, 2):
                    if not any(c.value is not None for c in cells):
                        continue
                    values = []
                    for index, cell in enumerate(cells):
                        if cell.value is None:
                            continue
                        value = cell.value
                        if isinstance(value, (date, datetime)):
                            value = value.isoformat()
                        values.append({"column": get_column_letter(index + 1),
                                       "header": headers[index], "value": value,
                                       "formula": cell.data_type == "f"})
                    yield {"file": path.name, "sheet": sheet.title, "row": number,
                           "locator": f"{path.name}::{sheet.title}::{number}", "cells": values}
        finally:
            book.close()


def one(raw, header):
    values = [c["value"] for c in raw["cells"]
              if norm(c["header"]) == norm(header) and meaningful(c["value"])]
    if len({str(v) for v in values}) > 1:
        raise ValueError(f"Valores conflitantes para {header}")
    return values[0] if values else None


def aliases(config, key):
    return {norm(k): v for k, v in config.get(key, {}).items()}


def classify(header, config, service_rule=None):
    key = norm(header)
    overrides = aliases(config, "header_overrides")
    overrides.update({norm(k): v for k, v in (service_rule or {}).get("header_overrides", {}).items()})
    if key in overrides:
        role = overrides[key]
        if role not in {"business", "region", "location", "description", "extra", "ignore"}:
            raise ValueError(f"Papel inválido para {header}: {role}")
        return role
    for role, names in config["header_roles"].items():
        if key in map(norm, names):
            return role
    for role, patterns in config.get("header_patterns", {}).items():
        if any(re.search(pattern, key) for pattern in patterns):
            return role
    return "extra"


def resolve_location(cells, config, catalog):
    """Intersect constraints, preferring detailed locations over regional fallback.

    Unknown details may create a location only under an unambiguous parent.
    A region-only source can use the explicitly catalogued unspecified location.
    """
    evidence = [c for c in cells if c["role"] in {"business", "region", "location"}
                and meaningful(c["value"])]
    if not evidence:
        fallback = config.get("missing_location")
        if fallback:
            if not any(all(norm(p[k]) == norm(fallback[k]) for k in ("business", "region", "location")) for p in catalog):
                return None, ["Localização padrão ausente do catálogo"]
            return dict(fallback), []
        return None, ["Localização ausente"]
    # Floor is a refinement of the building, not an independent global location.
    floors = [c for c in evidence if norm(c["header"]) in {"andar", "pavimento"}]
    business_map = aliases(config, "business_aliases")
    region_map = aliases(config, "region_aliases")
    location_map = aliases(config, "location_aliases")
    candidates = list(catalog)
    details = False
    errors = []
    unknown = []
    explicit_business = any(c["role"] == "business" for c in evidence)
    if not explicit_business and config.get("default_business"):
        candidates = [p for p in candidates if norm(p["business"]) == norm(config["default_business"])]
    for cell in evidence:
        if cell in floors:
            continue
        value = str(cell["value"]).strip()
        role = cell["role"]
        if role == "business":
            constraint = {"business": business_map.get(norm(value), value)}
        elif role == "region":
            constraint = {"region": region_map.get(norm(value), value)}
        else:
            mapped = location_map.get(norm(value), value)
            if isinstance(mapped, dict):
                constraint = mapped
                if not constraint or not set(constraint) <= {"business", "region", "location"}:
                    raise ValueError(f"Alias de localização inválido: {value}")
                details |= "location" in constraint
            elif any(norm(p["location"]) == norm(mapped) for p in catalog):
                constraint = {"location": mapped}
                details = True
            elif any(norm(p["region"]) == norm(region_map.get(norm(value), value)) for p in catalog):
                constraint = {"region": region_map.get(norm(value), value)}
            else:
                if config.get("create_unmapped_locations"):
                    unknown.append(str(mapped).strip())
                else:
                    errors.append(f"Local não mapeado ({cell['header']}): {value}")
                continue
        matches = [p for p in catalog if all(norm(p[k]) == norm(v) for k, v in constraint.items())]
        if not matches:
            errors.append(f"Localização fora do catálogo ({cell['header']}): {value}")
        candidates = [p for p in candidates if p in matches]
    if unknown:
        names = {norm(value): value for value in unknown}
        if len(names) > 1 or details or floors:
            return None, errors + ["Localização conflitante: novo local acompanhado de outro detalhe; conciliar os campos"]
        value = next(iter(names.values()))
        if not norm(value) or len(value) > 100:
            return None, errors + ["Nome do novo local inválido ou maior que 100 caracteres"]
        # Use explicit business/region and whole catalog names contained in the text.
        # No geographical guesses, fuzzy matches, or missing-location fallback here.
        padded = f" {norm(value)} "
        for key, name_map in (("business", business_map), ("region", region_map)):
            known = {norm(p[key]): p[key] for p in catalog}
            known.update(name_map)
            mentioned = {norm(canonical) for alias, canonical in known.items()
                         if f" {alias} " in padded}
            if mentioned:
                candidates = [p for p in candidates if norm(p[key]) in mentioned]
        # A known landmark (e.g. Refeitório) can also identify its parent region.
        landmarks = [p for p in catalog if norm(p["location"]) != "local exato nao especificado"
                     and f" {norm(p['location'])} " in padded]
        if landmarks:
            longest = max(len(norm(p["location"])) for p in landmarks)
            parents = {(norm(p["business"]), norm(p["region"])) for p in landmarks
                       if len(norm(p["location"])) == longest}
            candidates = [p for p in candidates if (norm(p["business"]), norm(p["region"])) in parents]
        parents = {(norm(p["business"]), norm(p["region"])): p for p in candidates}
        if errors or not parents:
            return None, errors + ["Localização conflitante: região do novo local incompatível com os dados"]
        if len(parents) != 1:
            return None, [f"Região/unidade do novo local não determinada: {value}"]
        parent = next(iter(parents.values()))
        return {"business": parent["business"], "region": parent["region"], "location": value}, []
    for cell in floors:
        floor = norm(cell["value"])
        number = re.search(r"\b(\d+)(?:o|a)?\b", floor)
        if number:
            pattern = rf"\b{int(number.group(1))}(?:o|a)?\s+(?:andar|pavimento)\b"
        elif floor in {"terreo", "andar terreo", "pavimento terreo"}:
            pattern = r"\bterreo\b"
        else:
            errors.append(f"Andar não mapeado: {cell['value']}")
            continue
        candidates = [p for p in candidates if re.search(pattern, norm(p["location"]))]
        if not details:
            generic_floor = [p for p in candidates if re.search(pattern + "$", norm(p["location"]))]
            if generic_floor:
                candidates = generic_floor
        details = True
    if not details:
        candidates = [p for p in candidates if norm(p["location"]) == "local exato nao especificado"]
    unique = {tuple(p[k] for k in ("business", "region", "location")): p for p in candidates}
    if len(unique) != 1:
        errors.append("Localização ambígua" if unique else "Localização conflitante ou não resolvida")
    return (next(iter(unique.values())) if len(unique) == 1 and not errors else None), errors


def parse_date(value):
    if not meaningful(value):
        return None
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    for fmt in ("%d/%m/%Y %H:%M", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(value).strip(), fmt).isoformat(timespec="seconds")
        except ValueError:
            pass
    raise ValueError(f"Data inválida: {value}")


BASE_HEADERS = {
    norm(x) for x in ("Número", "Categoria", "SubCategoria", "Tipo Chamado", "Chamado",
                      "Data Abertura", "Data Encerramento", "Histórico", "Data Agendamento", "Status",
                      "Nome Criador", "E-mail Criador", "Nome Solicitante", "E-mail Solicitante",
                      "Notas", "QrCode Ambiente", "SLA", "Reservas Vinculadas", "Data Reprovação",
                      "Data Cancelamento", "Nome Aprovador", "E-mail Aprovador", "Nome Cancelador",
                      "E-mail Cancelador")
}


def map_row(raw, config, catalog):
    errors, warnings = [], []
    result = {"raw": raw, "errors": errors, "warnings": warnings, "request": None}
    if any(c.get("formula") for c in raw["cells"]):
        errors.append("Planilha contém fórmula; exporte valores literais antes de importar")
    try:
        legacy = one(raw, "Número")
        if legacy is None:
            raise ValueError("Número legado ausente")
        legacy = str(int(legacy)) if isinstance(legacy, (int, float)) and int(legacy) == legacy else str(legacy).strip()
        override = config.get("row_overrides", {}).get(raw["locator"], {})
        legacy = str(override.get("legacy_id", legacy))
        if not legacy or len(legacy) > 100:
            raise ValueError("Identificador legado deve ter entre 1 e 100 caracteres")
        result["legacy_id"] = legacy
        category, service = one(raw, "Categoria"), one(raw, "Chamado")
        category = aliases(config, "category_aliases").get(norm(category), category)
        service_rule = {}
        for rule in config.get("service_rules", []):
            if norm(rule["category"]) == norm(category) and norm(rule["service"]) == norm(service):
                service_rule = rule
                break
        local_config = dict(config)
        local_config.update({k: v for k, v in service_rule.items() if k == "default_business"})
        cells = [dict(c, role=classify(c["header"], config, service_rule))
                 for c in raw["cells"] if meaningful(c["value"])]
        location, location_errors = resolve_location(cells, local_config, catalog)
        approved = config.get("location_resolutions", {}).get(location_signature(cells))
        if location_errors and approved:
            location = dict(approved["location"])
            location_errors = []
            warnings.append("Localização definida pela regra revisada: " + approved["reason"])
        if "location" in override:
            location = override["location"]
            location_errors = [] if location in catalog else ["Localização do override fora do catálogo"]
        errors.extend(location_errors)
        if location and not any(c["role"] in {"business", "region", "location"} for c in cells):
            warnings.append("Localização ausente: aplicada Região 1 / Prédio Administrativo / Local exato não especificado")
        elif location and not any(all(norm(location[k]) == norm(p[k]) for k in ("business", "region", "location")) for p in catalog):
            warnings.append("Novo local: cadastrar ou reutilizar pelo nome na região resolvida")
        descriptions = []
        for cell in cells:
            if cell["role"] == "description":
                value = str(cell["value"]).strip()
                if value not in descriptions:
                    descriptions.append(value)
        description = override.get("description", config.get("description_separator", "\n").join(descriptions) or None)
        if description and len(description) > 300:
            summary = config.get("description_summaries", {}).get(legacy)
            if summary and summary["source_hash"] == digest(description):
                result["full_description"] = description
                description = summary["text"]
                warnings.append("Descrição resumida; texto integral preservado no snapshot")
        if description and len(description) > 300:
            errors.append(f"Descrição com {len(description)} caracteres (limite 300); definir resumo em row_overrides")
        email = str(one(raw, "E-mail Solicitante") or "").strip().lower()
        if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", email):
            errors.append("E-mail do solicitante ausente ou inválido")
        history = str(one(raw, "Histórico") or "")
        created = parse_date(one(raw, "Data Abertura"))
        if not created:
            dates = re.findall(r"Data abertura:\s*(\d{2}/\d{2}/\d{4} \d{2}:\d{2})", history)
            if dates:
                created = min(parse_date(d) for d in dates)
                warnings.append("Abertura recuperada do primeiro evento explícito do histórico")
            else:
                warnings.append("Data de abertura ausente na origem e no histórico")
        started = re.findall(r"Data andamento:\s*(\d{2}/\d{2}/\d{4} \d{2}:\d{2})", history)
        status = aliases(config, "status_map").get(norm(one(raw, "Status")))
        request_type = aliases(config, "request_type_map").get(norm(one(raw, "Tipo Chamado")))
        if status not in (1, 2, 3, 4, 5):
            errors.append("Status não mapeado")
        if request_type not in (1, 2):
            errors.append("Tipo de chamado não mapeado")
        for name, value, maximum in (("Categoria", category, 100), ("Chamado", service, 100),
                                     ("SubCategoria", one(raw, "SubCategoria"), 100),
                                     ("Solicitante", one(raw, "Nome Solicitante"), 300), ("E-mail", email, 300)):
            if value and len(str(value)) > maximum:
                errors.append(f"{name} excede {maximum} caracteres")
        if not category or not service:
            errors.append("Categoria ou serviço ausente")
        fields = []
        field_map = aliases(config, "field_aliases")
        field_map.update({norm(k): v for k, v in service_rule.get("field_aliases", {}).items()})
        for cell in cells:
            if cell["role"] != "extra" or norm(cell["header"]) in BASE_HEADERS or norm(cell["header"]).startswith("sla "):
                continue
            if re.search(r"https?://", str(cell["value"]), re.I):
                from .media import urls
                links = urls(cell["value"])
                if not links:
                    continue
            else:
                links = None
            name = field_map.get(norm(cell["header"]), cell["header"].strip().rstrip(":"))
            if len(name) > 100 and links:
                candidates = [d["name"] for d in config.get("field_catalog", [])
                              if norm(d["category"]) == norm(category) and norm(d["service"]) == norm(service)
                              and len(d["name"]) >= 80 and norm(name).startswith(norm(d["name"]))]
                if len(set(candidates)) == 1:
                    name = candidates[0]
            if len(name) > 100:
                errors.append(f"Campo maior que 100 caracteres; definir field_aliases: {cell['header']}")
            fields.append({"name": name, "value": links if links else cell["value"],
                           **({"media": True} if links else {})})
        by_name = defaultdict(list)
        for field in fields:
            by_name[norm(field["name"])].append(field)
        if any(len({digest(f["value"]) for f in group}) > 1 for group in by_name.values()):
            errors.append("Respostas conflitantes para o mesmo campo adicional")
        mapped_fields = [group[0] for group in by_name.values()]
        if config.get("field_catalog") or any(f.get("media") for f in mapped_fields):
            from .field_catalog import prepare_fields
            prepare_fields(mapped_fields, category, service, config, warnings)
        result["request"] = {
            "category": category, "service": service, "subcategory": one(raw, "SubCategoria"),
            "request_type_id": request_type, "status_id": status,
            "requester": {"email": email, "name": one(raw, "Nome Solicitante")},
            "location": location, "description": description,
            "created_date": created, "agreed_date": parse_date(one(raw, "Data Agendamento")),
            "started_date": min(map(parse_date, started)) if started else None,
            "finished_date": parse_date(one(raw, "Data Encerramento")),
            "canceled_date": parse_date(one(raw, "Data Cancelamento")),
            "fields": mapped_fields,
            "existing_request_id": config.get("existing_request_ids", {}).get(legacy),
        }
        occurrences = config.get("duplicate_occurrences", {}).get(legacy)
        result["original_number"] = legacy
        if occurrences:
            matched = [rule for rule in occurrences if all(result["request"].get(k) == v for k, v in rule["match"].items())]
            if len(matched) != 1:
                errors.append("Ocorrência do número repetido não identificada pelas regras revisadas")
            else:
                result["legacy_id"] = f"{legacy}#{matched[0]['occurrence']}"
        if "request_numbers" in config:
            number = config["request_numbers"].get(result["legacy_id"])
            if number is None:
                errors.append("Número de destino não reservado; executar --reserve-numbers")
            elif type(number) is not int or not 0 < number < 10**10:
                errors.append("Número de destino inválido")
            else:
                result["request"]["import_id"] = number
        # Location/date/description semantics are included; filename and row are not identity.
        original = sorted((norm(c["header"]), str(c["value"])) for c in raw["cells"])
        result["content_hash"] = digest({"request": result["request"], "original": original})
    except (ValueError, TypeError, KeyError) as exc:
        errors.append(str(exc))
    return result


def build_plan(raw_rows, config, catalog):
    if not isinstance(config.get("source"), str) or not 0 < len(config["source"]) <= 100:
        raise ValueError("source deve identificar o sistema/tenant de origem em até 100 caracteres")
    raw_rows = list(raw_rows)
    excluded_categories = {norm(c) for c in config.get("excluded_categories", [])}
    category_aliases = aliases(config, "category_aliases")
    included, excluded = [], []
    for raw in raw_rows:
        try:
            category = one(raw, "Categoria")
        except ValueError:
            included.append(raw)  # map_row reports conflicting source cells as pending.
            continue
        category = category_aliases.get(norm(category), category)
        (excluded if norm(category) in excluded_categories else included).append(raw)
    rows = [map_row(raw, config, catalog) for raw in included]
    groups = defaultdict(list)
    for row in rows:
        if "legacy_id" in row:
            groups[row["legacy_id"]].append(row)
    duplicates = 0
    for group in groups.values():
        if len(group) <= 1:
            continue
        if len({r.get("content_hash") for r in group}) == 1 and all(not r["errors"] for r in group):
            for row in group[1:]:
                row["duplicate_of"] = group[0]["raw"]["locator"]
                duplicates += 1
        else:
            for row in group:
                row["errors"].append("Número legado repetido com conteúdo conflitante; conciliar antes da carga")
    pending = sum(bool(r["errors"]) for r in rows)
    numbers = defaultdict(list)
    for row in rows:
        if row["request"] and row["request"].get("import_id"):
            numbers[row["request"]["import_id"]].append(row)
    for group in numbers.values():
        if len({r["legacy_id"] for r in group}) > 1:
            for row in group:
                row["errors"].append("Número de destino reservado para mais de uma origem")
    pending = sum(bool(r["errors"]) for r in rows)
    new_locations = {}
    for row in rows:
        if not row["request"] or not row["request"]["location"]:
            continue
        location = row["request"]["location"]
        key = tuple(norm(location[k]) for k in ("business", "region", "location"))
        if not any(key == tuple(norm(p[k]) for k in ("business", "region", "location")) for p in catalog):
            new_locations.setdefault(key, location)
    return {"version": 1, "source": config["source"], "mapping_hash": digest({"config": config, "catalog": catalog}),
            "excluded_categories": config.get("excluded_categories", []), "excluded_rows": excluded,
            "new_locations": list(new_locations.values()),
            "summary": {"rows": len(raw_rows), "excluded": len(excluded), "ready": len(rows) - pending - duplicates,
                        "pending": pending, "identical_duplicates": duplicates,
                        "warnings": sum(len(r["warnings"]) for r in rows)},
            "pending_reasons": dict(Counter(e for r in rows for e in r["errors"])), "rows": rows}


def reserve_numbers(rows, existing):
    """Keep prior reservations immutable; append fresh keys without ID collisions."""
    numbers = dict(existing)
    if any(type(v) is not int or not 0 < v < 10**10 for v in numbers.values()) or len(set(numbers.values())) != len(numbers):
        raise ValueError("Reservas de numeração inválidas ou duplicadas")
    keys = {r["legacy_id"] for r in rows if r.get("request") and r.get("legacy_id")
            and not any("Ocorrência" in e or "Número legado repetido" in e for e in r["errors"])}
    def order(key):
        parts = key.split("#")
        return int(parts[0]), int(parts[1]) if len(parts) > 1 else 0
    used = set(numbers.values())
    fresh = sorted(keys - numbers.keys(), key=order)
    # Reserve the approved adjacent duplicate pairs before relocating ordinary
    # records that formerly occupied their second numbers. Existing reservations
    # are immutable across later exports and partial batches.
    fresh = [k for k in fresh if "#" in k] + [k for k in fresh if "#" not in k]
    for key in fresh:
        number = order(key)[0] + (order(key)[1] - 1 if "#" in key else 0)
        if "#" in key and number in used:
            raise ValueError(f"Número da ocorrência {key} já reservado; conciliar a numeração existente")
        while number in used:
            number += 1
        if number >= 10**10:
            raise ValueError("Numeração excede NUMBER(10)")
        numbers[key] = number
        used.add(number)
    return numbers


def render_report(plan):
    summary = plan["summary"]
    lines = ["# Prévia da importação", "", f"Origem: {plan['source']}", "",
             f"- Linhas: {summary['rows']}", f"- Resolvidas pelas regras: {summary['ready']}",
             f"- Pendentes: {summary['pending']}",
             f"- Excluídas por categoria: {summary.get('excluded', 0)}",
             f"- Repetições idênticas: {summary['identical_duplicates']}", "",
             "A prévia não consulta o banco. Cadastros ambíguos, tipos de campos e correspondências existentes são validados na aplicação.",
             "Anexos externos ficam preservados como URLs na origem, sem download para BLOB.", "",
             "## Novos locais previstos", "",
             "Criação ou reutilização durante --apply, somente para linhas importadas. Sem gravação na prévia.", ""]
    for location in plan.get("new_locations", []):
        lines.append(f"- {location['business']} / {location['region']} / {location['location']}")
    lines.extend(["", "## Números de origem e destino alterados", "",
                  "As reservas são persistentes; arquivo original e chave de origem permanecem no snapshot.", ""])
    for row in plan["rows"]:
        if row.get("request") and row["request"].get("import_id") and str(row["request"]["import_id"]) != row.get("original_number"):
            lines.append(f"- Origem {row['legacy_id']} → REQUEST.ID {row['request']['import_id']}")
    lines.extend(["", "## Descrições resumidas", ""])
    for row in plan["rows"]:
        if row.get("full_description") and row.get("request"):
            lines.append(f"- Origem {row['legacy_id']}: {row['request']['description']}")
    lines.extend(["", "## Motivos de pendência", "", "Uma linha pode ter mais de um motivo.", ""])
    for reason, count in sorted(plan["pending_reasons"].items(), key=lambda item: (-item[1], item[0])):
        lines.append(f"- {count}: {reason}")
    lines.extend(["", "## Linhas pendentes", ""])
    for row in plan["rows"]:
        if row["errors"]:
            lines.append(f"- {row.get('legacy_id', 'Sem número')} — {row['raw']['locator']}: " + "; ".join(row["errors"]))
    return "\n".join(lines) + "\n"
